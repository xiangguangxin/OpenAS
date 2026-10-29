#!/usr/bin/env python3
"""Local Office operations. No installation, source overwrite, or remote converter.

export-pdf checks output signature, render checks PDF/page-image structure, and
recalculate checks preserved formulas and cache availability. None certifies
visual quality, business logic, or native Microsoft Office compatibility.
"""
import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import zipfile

OFFICE_TYPES = {'.docx', '.xlsx', '.pptx'}


def executable(name, default):
    path = shutil.which(str(name)) if name else next((p for item in default if (p := shutil.which(item))), None)
    if not path:
        raise ValueError('Required local tool unavailable: ' + (str(name) if name else '/'.join(default)))
    return path


def run_command(argv, directory, timeout):
    """Keep diagnostics and await process termination after timeout/interruption."""
    with (directory/'operation.log').open('a', encoding='utf-8') as log:
        process = subprocess.Popen(argv, stdout=log, stderr=log, start_new_session=True)
        try:
            code = process.wait(timeout=timeout)
        except (subprocess.TimeoutExpired, KeyboardInterrupt):
            if os.name == 'posix':
                os.killpg(process.pid, signal.SIGTERM)
            else:
                process.terminate()
            try:
                process.wait(timeout=3)
            except subprocess.TimeoutExpired:
                if os.name == 'posix':
                    os.killpg(process.pid, signal.SIGKILL)
                else:
                    process.kill()
                process.wait()
            raise
    if code:
        raise ValueError(f'Local command failed ({code}); see operation.log')


def formula_map(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=False)
    try:
        result = {}
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.data_type == 'f':
                        if not isinstance(cell.value, str):
                            raise ValueError('Array/data-table formula preservation requires a dedicated workflow')
                        result[(sheet.title, cell.coordinate)] = cell.value
        return result
    finally:
        workbook.close()


def verify_recalculation(original, recalculated):
    from openpyxl import load_workbook
    formulas = formula_map(original)
    if formula_map(recalculated) != formulas:
        raise ValueError('Formula locations or expressions changed during recalculation')
    workbook = load_workbook(recalculated, data_only=True)
    try:
        missing, errors = [], []
        for (sheet, address) in formulas:
            cell = workbook[sheet][address]
            label = f'{sheet}!{address}'
            if cell.data_type == 'e':
                errors.append(label + ': ' + str(cell.value))
            elif cell.value is None and cell.data_type != 'str':
                missing.append(label)
        return {'formula_cells': len(formulas), 'formulas_preserved': True,
                'missing_cache_cells': missing, 'formula_errors': errors,
                'business_logic_review': 'not_run'}
    finally:
        workbook.close()


def page_selection(specification, count):
    if specification is None:
        return list(range(1, count + 1))
    try:
        pages = [int(part.strip()) for part in specification.split(',')]
    except (ValueError, AttributeError) as exc:
        raise ValueError('pages must be comma-separated 1-based page numbers') from exc
    if not pages or len(set(pages)) != len(pages) or any(p < 1 or p > count for p in pages):
        raise ValueError('Page selection is empty, duplicated or outside the PDF')
    return pages


def perform(operation, source, output, *, soffice=None, pdftoppm=None, pages=None, timeout=120):
    source, output = Path(source).resolve(), Path(output).resolve()
    if operation not in ('export-pdf', 'render', 'recalculate'):
        raise ValueError('Unknown Office operation')
    allowed = {'.xlsx'} if operation == 'recalculate' else OFFICE_TYPES | ({'.pdf'} if operation == 'render' else set())
    if not source.is_file() or source.suffix.lower() not in allowed or not math.isfinite(timeout) or timeout <= 0:
        raise ValueError('Existing supported source file and positive timeout required')
    if output.exists():
        raise ValueError('Output directory must be new; source and existing outputs are preserved')
    if pages is not None and operation != 'render':
        raise ValueError('Page selection applies only to render')
    office = executable(soffice, ('soffice', 'libreoffice')) if source.suffix.lower() != '.pdf' else None
    poppler = executable(pdftoppm, ('pdftoppm',)) if operation == 'render' else None
    if operation == 'render':
        from pypdf import PdfReader
    if operation == 'recalculate':
        # Conversion is limited to local ordinary XLSX workbooks, not external refresh.
        with zipfile.ZipFile(source) as package:
            if any(name.startswith(('xl/externalLinks/', 'xl/queryTables/')) or name in
                   ('xl/connections.xml', 'xl/vbaProject.bin') for name in package.namelist()):
                raise ValueError('External-data or macro workbooks require a dedicated recalculation workflow')
        for formula in formula_map(source).values():
            if re.search(r'\b(?:WEBSERVICE|RTD|DDE|IMAGE)\s*\(|\[[^\]]+\][^!]*!', formula, re.I):
                raise ValueError('External-data formulas require a dedicated recalculation workflow')
    source_bytes = source.read_bytes()
    report = {'version': 1, 'operation': operation, 'source_name': source.name,
              'source_sha256': hashlib.sha256(source_bytes).hexdigest(), 'status': 'running',
              'visual_review': 'not_run', 'outputs': [],
              'limits': 'Local conversion is not visual, business-logic, or native Office acceptance.'}
    output.mkdir(parents=True, exist_ok=False)
    deadline = time.monotonic() + timeout
    def run(argv):
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise subprocess.TimeoutExpired(argv, timeout)
        run_command(argv, output, remaining)
    try:
        with tempfile.TemporaryDirectory(prefix='openas-office-') as temporary:
            temp = Path(temporary)
            staged = temp/('source' + source.suffix.lower())
            staged.write_bytes(source_bytes)
            if operation == 'recalculate':
                from openpyxl import load_workbook
                from openpyxl.workbook.properties import CalcProperties
                original_snapshot = temp/'original.xlsx'
                original_snapshot.write_bytes(source_bytes)
                workbook = load_workbook(staged, data_only=False)
                try:
                    # Clear old formula caches before invoking the real engine.
                    # A mere copy/resave without calculation must fail verification.
                    workbook.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
                    workbook.save(staged)
                finally:
                    workbook.close()
            if office:
                converted = temp/'converted'; converted.mkdir()
                profile = temp/'profile'; (profile/'user').mkdir(parents=True)
                # New profile with no trusted macro locations; do not use the user's profile.
                (profile/'user/registrymodifications.xcu').write_text(
                    '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                    '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
                    '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>'
                    '</item></oor:items>', encoding='utf-8')
                format_ = 'xlsx:Calc MS Excel 2007 XML' if operation == 'recalculate' else 'pdf'
                run([office, '-env:UserInstallation=' + profile.as_uri(), '--headless',
                     '--convert-to', format_, '--outdir', str(converted), str(staged)])
                artifact = converted/('source.xlsx' if operation == 'recalculate' else 'source.pdf')
            else:
                artifact = staged
            target = output/('recalculated.xlsx' if operation == 'recalculate' else 'document.pdf')
            if not artifact.is_file() or not artifact.stat().st_size:
                raise ValueError('Converter returned without the expected nonempty artifact')
            shutil.copyfile(artifact, target)
            report['outputs'].append(target.name)
            if operation == 'recalculate':
                report['recalculation'] = verify_recalculation(original_snapshot, target)
                checks = report['recalculation']
                if checks['missing_cache_cells'] or checks['formula_errors']:
                    raise ValueError('Formula caches are missing or contain errors; inspect the report')
            else:
                if not target.read_bytes().startswith(b'%PDF-'):
                    raise ValueError('Output has no PDF signature')
                report['pdf_validation'] = 'nonempty PDF signature only'
                if operation == 'render':
                    reader = PdfReader(target)
                    if reader.is_encrypted:
                        raise ValueError('Encrypted PDF requires an explicitly supported workflow')
                    count = len(reader.pages)
                    if count < 1:
                        raise ValueError('PDF has no pages')
                    selected = page_selection(pages, count)
                    directory = output/'pages'; directory.mkdir()
                    for number in selected:
                        prefix = directory/f'page-{number:04d}'
                        run([poppler, '-f', str(number), '-l', str(number), '-singlefile',
                             '-png', '-scale-to', '1600', str(target), str(prefix)])
                        image = prefix.with_suffix('.png')
                        if not image.is_file() or not image.read_bytes().startswith(b'\x89PNG\r\n\x1a\n'):
                            raise ValueError('Renderer returned without the expected PNG')
                        report['outputs'].append(str(image.relative_to(output)))
                    report.update(pdf_page_count=count, rendered_pages=selected)
                    report['pdf_validation'] = 'PDF parsed; selected PNG signatures checked'
            report['output_sha256'] = {name: hashlib.sha256((output/name).read_bytes()).hexdigest()
                                       for name in report['outputs']}
            report['status'] = 'completed'
    except (Exception, KeyboardInterrupt) as exc:
        report.update(status='failed', error=str(exc) or type(exc).__name__)
        raise
    finally:
        (output/'operation-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('operation', choices=['export-pdf', 'render', 'recalculate'])
    parser.add_argument('source', type=Path)
    parser.add_argument('--output-dir', required=True, type=Path)
    parser.add_argument('--soffice')
    parser.add_argument('--pdftoppm')
    parser.add_argument('--pages', help='render only: comma-separated 1-based PDF page numbers')
    parser.add_argument('--timeout', type=float, default=120)
    args = parser.parse_args()
    try:
        report = perform(args.operation, args.source, args.output_dir, soffice=args.soffice,
                         pdftoppm=args.pdftoppm, pages=args.pages, timeout=args.timeout)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        print(f'Office operation failed: {exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.exit(main())
