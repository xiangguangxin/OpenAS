#!/usr/bin/env python3
"""Copyright 2026 OpenAS contributors. SPDX-License-Identifier: Apache-2.0.
Build editable slides from a bounded JSON source. No network or dependency installation.
"""
import argparse
from datetime import datetime, timezone
import importlib.util
import io
import json
from pathlib import Path
import platform
import shutil
import sys

sys.dont_write_bytecode = True
import deck_model as model

DEFAULT_THEME = {'background': 'F7F8FA', 'foreground': '18232F', 'accent': '007D78',
                 'surface': 'E6EFED', 'muted': '516270', 'font': 'Arial', 'cjk_font': 'Microsoft YaHei'}


def environment():
    return {'python': platform.python_version(), 'python_pptx': importlib.util.find_spec('pptx') is not None,
            'pillow': importlib.util.find_spec('PIL') is not None,
            'soffice': shutil.which('soffice') or shutil.which('libreoffice'), 'pdftoppm': shutil.which('pdftoppm'),
            'font_check': 'Choose installed Latin/CJK fonts and verify them in actual rendered output; no fonts are bundled'}


def build(spec, assets_root):
    model.validate(spec, assets_root)
    try:
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise ValueError('Install the reviewed requirements.txt into a task-local virtual environment') from exc
    theme = {**DEFAULT_THEME, **spec.get('theme', {})}
    prs = Presentation()
    width, height = model.SIZES[spec.get('size', 'wide')]
    prs.slide_width, prs.slide_height = Inches(width), Inches(height)
    prs.core_properties.title = spec['title']
    prs.core_properties.subject = 'Editable presentation generated from an OpenAS JSON specification'
    prs.core_properties.author = ''
    prs.core_properties.last_modified_by = 'OpenAS'
    prs.core_properties.created = prs.core_properties.modified = datetime.now(timezone.utc)
    margin, usable = .65, width-1.3

    def font(run, size, color, bold=False):
        run.font.name = theme['font']; run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        run.font.bold = bold
        props = run._r.get_or_add_rPr()
        ea = OxmlElement('a:ea'); ea.set('typeface', theme['cjk_font']); props.append(ea)

    def textbox(slide, name, content, x, y, w, h, size=24, color=None, bold=False):
        if model.estimated_lines(content, w-.06, size)*size*1.16/72 > h-.04:
            raise ValueError(f'{name}: content exceeds conservative text budget; shorten or split the slide')
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.name = name
        frame = box.text_frame
        frame.clear(); frame.word_wrap = True; frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = Inches(.02)
        frame.margin_top = frame.margin_bottom = Inches(.01)
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for index, line in enumerate(content.split('\n')):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.space_before = Pt(0); paragraph.space_after = Pt(0); paragraph.line_spacing = 1.16
            run = paragraph.add_run(); run.text = line
            font(run, size, color or theme['foreground'], bold)
        return box

    def shape(slide, name, x, y, w, h, color):
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.name = 'decor:' + name
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor.from_string(color); box.line.fill.background()
        return box

    def bullets(slide, items, x, y, w, h, size=24, name='body'):
        rows = [model.estimated_lines(item, w-.35, size)*size*1.16/72+.18 for item in items]
        if sum(rows) > h:
            raise ValueError(f'{name}: bullet content is too dense; split the slide')
        for index, (item, row) in enumerate(zip(items, rows)):
            shape(slide, f'{name}-marker-{index}', x, y+.12, .075, .075, theme['accent'])
            textbox(slide, f'{name}-{index}', item, x+.3, y, w-.3, row-.12, size)
            y += row

    for number, item in enumerate(spec['slides'], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(theme['background'])
        kind = item['type']
        if kind in ('cover', 'section'):
            shape(slide, 'cover-panel', 0, 0, width, height, theme['foreground'])
            textbox(slide, 'cover-kicker', spec.get('audience', 'PRESENTATION'), margin, .75, usable, .45, 16, theme['surface'])
            textbox(slide, 'title', item['title'], margin, 1.8, usable, 1.6, 38, 'FFFFFF', True)
            if item.get('subtitle'):
                textbox(slide, 'subtitle', item['subtitle'], margin, 3.8, usable, 1.35, 23, theme['surface'])
        else:
            textbox(slide, 'title', item['title'], margin, .55, usable, 1.05, 32, bold=True)
            top = 2.0
            if kind in ('bullets', 'closing'):
                bullets(slide, item['bullets'], margin, top, usable, 3.8 if item.get('caption') else 4.35, 25)
            elif kind == 'two-column':
                gap, col = .5, (usable-.5)/2
                for index, side in enumerate(('left', 'right')):
                    x = margin+index*(col+gap)
                    textbox(slide, side+'-heading', item[side]['heading'], x, top, col, .75, 24, theme['accent'], True)
                    bullets(slide, item[side]['bullets'], x, top+1.0, col, 2.8 if item.get('caption') else 3.35, 22, side)
            elif kind == 'process':
                gap = .24; col = (usable-gap*(len(item['steps'])-1))/len(item['steps'])
                for index, step in enumerate(item['steps']):
                    x = margin+index*(col+gap)
                    shape(slide, 'process-card', x, top+.1, col, 3.8, theme['surface'])
                    textbox(slide, f'step-number-{index}', f'{index+1:02}', x+.18, top+.32, col-.36, .7, 34, theme['accent'], True)
                    textbox(slide, f'step-label-{index}', step['label'], x+.18, top+1.3, col-.36, .9, 23, bold=True)
                    textbox(slide, f'step-detail-{index}', step['detail'], x+.18, top+2.35, col-.36, 1.25, 20)
            elif kind == 'chart':
                data = CategoryChartData(); data.categories = item['categories']
                for series in item['series']:
                    data.add_series(series['name'], series['values'])
                chart_kind = XL_CHART_TYPE.LINE if item.get('chart_type') == 'line' else XL_CHART_TYPE.COLUMN_CLUSTERED
                frame = slide.shapes.add_chart(chart_kind, Inches(margin), Inches(top), Inches(usable), Inches(3.9), data)
                frame.name = 'native-chart'; chart = frame.chart
                chart.chart_style = 10
                chart.has_title = False
                chart.has_legend = len(item['series']) > 1
                if chart.has_legend:
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    chart.legend.include_in_layout = False
                    chart.legend.font.size = Pt(14)
                    chart.legend.font.name = theme['cjk_font']
                chart.category_axis.tick_labels.font.size = Pt(15)
                chart.category_axis.tick_labels.font.name = theme['cjk_font']
                chart.value_axis.tick_labels.font.size = Pt(14)
                chart.value_axis.has_title = True
                chart.value_axis.axis_title.text_frame.text = item['unit']
                for paragraph in chart.value_axis.axis_title.text_frame.paragraphs:
                    for run in paragraph.runs:
                        font(run, 14, theme['muted'])
                if len(item['series']) == 1:
                    if chart_kind == XL_CHART_TYPE.LINE:
                        chart.series[0].format.line.color.rgb = RGBColor.from_string(theme['accent'])
                    else:
                        chart.series[0].format.fill.solid()
                        chart.series[0].format.fill.fore_color.rgb = RGBColor.from_string(theme['accent'])
                if chart_kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
                    chart.plots[0].has_data_labels = True
                    chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
                    chart.plots[0].data_labels.font.size = Pt(14)
                for props in chart._chartSpace.xpath('.//a:defRPr'):
                    ea = OxmlElement('a:ea'); ea.set('typeface', theme['cjk_font']); props.append(ea)
            elif kind == 'table':
                rows, cols = len(item['rows'])+1, len(item['columns'])
                frame = slide.shapes.add_table(rows, cols, Inches(margin), Inches(top), Inches(usable), Inches(3.9))
                frame.name = 'native-table'; table = frame.table
                for row_index, values in enumerate([item['columns'], *item['rows']]):
                    for col_index, value in enumerate(values):
                        cell = table.cell(row_index, col_index)
                        cell.margin_left = cell.margin_right = Inches(.12)
                        cell.margin_top = cell.margin_bottom = Inches(.06)
                        if model.estimated_lines(value, usable/cols-.24, 19)*19*1.16/72 > 3.9/rows-.12:
                            raise ValueError('Table cell exceeds text budget; simplify the table')
                        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(theme['foreground'] if row_index == 0 else theme['surface'])
                        paragraph = cell.text_frame.paragraphs[0]
                        run = paragraph.add_run(); run.text = value
                        font(run, 19, 'FFFFFF' if row_index == 0 else theme['foreground'], row_index == 0)
            elif kind == 'image':
                from PIL import Image
                path = model.image_path(assets_root, item['image'])
                with Image.open(path) as image:
                    image_w, image_h = image.size
                scale = min(usable/image_w, 3.9/image_h)
                w, h = image_w*scale, image_h*scale
                picture = slide.shapes.add_picture(str(path), Inches(margin+(usable-w)/2), Inches(top), Inches(w), Inches(h))
                picture.name = 'original-image'
                picture._element.nvPicPr.cNvPr.set('descr', item['alt'])
            if item.get('caption'):
                textbox(slide, 'caption', item['caption'], margin, 6.05, usable, .55, 15, theme['muted'])
        sources = item.get('sources', [])
        if sources:
            label = sources[0]+(f' (+{len(sources)-1})' if len(sources)>1 else '')
            textbox(slide, 'source', label, margin, 6.95, usable-.65, .38, 10, theme['surface'] if kind in ('cover','section') else theme['muted'])
        if kind != 'cover':
            textbox(slide, 'page-number', str(number), width-1.0, 6.95, .4, .32, 12, theme['surface'] if kind == 'section' else theme['muted'])
        slide.notes_slide.notes_text_frame.text = (item['notes'] + '\n\nSpeaking budget: ' + str(item['seconds']) + ' seconds'
                                                  + ('\n\nSources:\n'+'\n'.join(sources) if sources else ''))
    result = io.BytesIO(); prs.save(result)
    return result.getvalue()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check-environment', action='store_true')
    parser.add_argument('--spec', type=Path)
    parser.add_argument('--assets-root', type=Path)
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    try:
        if args.check_environment:
            report = environment(); print(json.dumps(report, ensure_ascii=False, indent=2))
            return 0 if report['python_pptx'] and report['pillow'] else 1
        if not args.spec or not args.output:
            parser.error('--spec and --output are required for generation')
        if args.output.suffix.lower() != '.pptx':
            raise ValueError('Output must have .pptx extension')
        assets_root = (args.assets_root or args.spec.parent).resolve()
        raw = build(json.loads(args.spec.read_text(encoding='utf-8')), assets_root)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        with args.output.open('xb') as stream:
            stream.write(raw)
        report = model.inspect_pptx(args.output)
        print(json.dumps({'output': str(args.output), 'slide_count': report['slide_count'],
                          'structure_passed': report['passed'], 'rendered_review': 'not_run'}, ensure_ascii=False))
        return 0 if report['passed'] else 1
    except (OSError, ValueError, TypeError, KeyError) as exc:
        print(f'Build failed: {exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.exit(main())
