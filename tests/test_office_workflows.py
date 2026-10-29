"""Office artifacts and operation boundaries; mocked conversions are not renderer acceptance."""
import importlib.util
import io
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]
SKILLS = ROOT/'plugins/openas/skills'
SPEC = importlib.util.spec_from_file_location('office_ops', SKILLS/'office/scripts/office_ops.py')
ops = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(ops)
HAS_OFFICE = all(importlib.util.find_spec(name) is not None for name in ('docx', 'openpyxl', 'pypdf'))


class OperationBoundaryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root/'original.docx'; self.source.write_bytes(b'fixture source')
        self.output = self.root/'output'

    def test_missing_engine_does_not_create_output_or_modify_source(self):
        with patch.object(ops.shutil, 'which', return_value=None):
            with self.assertRaisesRegex(ValueError, 'Required local tool'):
                ops.perform('export-pdf', self.source, self.output)
        self.assertFalse(self.output.exists())
        self.assertEqual(self.source.read_bytes(), b'fixture source')

    def test_existing_outputs_and_unsupported_inputs_are_preserved(self):
        self.output.mkdir(); (self.output/'keep').write_text('existing')
        for source, timeout in ((self.source, 1), (self.root/'missing.docx', 1),
                                (self.source, float('nan')), (self.source, 0)):
            with self.subTest(source=source, timeout=timeout), self.assertRaises(ValueError):
                ops.perform('export-pdf', source, self.output, timeout=timeout)
        self.assertEqual((self.output/'keep').read_text(), 'existing')

    def test_export_uses_isolated_copy_and_retains_original(self):
        def convert(argv, directory, timeout):
            staged = Path(argv[-1])
            self.assertNotEqual(staged, self.source)
            self.assertEqual(staged.read_bytes(), self.source.read_bytes())
            self.assertIn('-env:UserInstallation=', argv[1])
            (Path(argv[argv.index('--outdir')+1])/'source.pdf').write_bytes(b'%PDF-1.4\nfixture')
        with patch.object(ops, 'executable', return_value='/fixture/soffice'), patch.object(ops, 'run_command', side_effect=convert):
            report = ops.perform('export-pdf', self.source, self.output)
        self.assertEqual(report['status'], 'completed')
        self.assertEqual(report['outputs'], ['document.pdf'])
        self.assertEqual(report['visual_review'], 'not_run')
        self.assertEqual(self.source.read_bytes(), b'fixture source')
        self.assertEqual(list(self.output.glob('pages/*')), [])

    def test_successful_process_without_artifact_is_not_success(self):
        with patch.object(ops, 'executable', return_value='/fixture/soffice'), patch.object(ops, 'run_command'):
            with self.assertRaisesRegex(ValueError, 'expected nonempty artifact'):
                ops.perform('export-pdf', self.source, self.output)
        report = json.loads((self.output/'operation-report.json').read_text())
        self.assertEqual(report['status'], 'failed')
        self.assertEqual(self.source.read_bytes(), b'fixture source')

    def test_non_pdf_artifact_fails_validation(self):
        def convert(argv, directory, timeout):
            (Path(argv[argv.index('--outdir')+1])/'source.pdf').write_text('not a PDF')
        with patch.object(ops, 'executable', return_value='/fixture/soffice'), patch.object(ops, 'run_command', side_effect=convert):
            with self.assertRaisesRegex(ValueError, 'PDF signature'):
                ops.perform('export-pdf', self.source, self.output)
        self.assertEqual(json.loads((self.output/'operation-report.json').read_text())['status'], 'failed')

    def test_timeout_waits_for_owned_process_exit(self):
        self.output.mkdir()
        with self.assertRaises(subprocess.TimeoutExpired):
            ops.run_command([sys.executable, '-c', 'import time; time.sleep(30)'], self.output, .1)
        self.assertTrue((self.output/'operation.log').is_file())

    def test_page_selection_rejects_out_of_range_and_duplicates(self):
        self.assertEqual(ops.page_selection('3,1', 3), [3, 1])
        for value in ('0', '4', '1,1', '', '1-3'):
            with self.subTest(value=value), self.assertRaises(ValueError):
                ops.page_selection(value, 3)


@unittest.skipUnless(HAS_OFFICE, 'Optional office/requirements.txt dependencies required')
class OfficeArtifactTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def run_examples(self, name):
        text = (SKILLS/name/'references/authoring.md').read_text()
        for block in re.findall(r'```python\n(.*?)\n```', text, re.S):
            subprocess.run([sys.executable, '-B', '-c', block], cwd=self.root, check=True, capture_output=True)

    def workbook(self):
        from openpyxl import Workbook
        source = self.root/'source.xlsx'
        book = Workbook(); sheet = book.active; sheet.title = 'Revenue'
        sheet.append([12]); sheet.append([18]); sheet.append(['=SUM(A1:A2)'])
        book.save(source); book.close()
        return source

    def cached_copy(self, source, value='30', kind='n'):
        output = self.root/'engine-result.xlsx'
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, 'w') as target:
            for item in original.infolist():
                data = original.read(item.filename)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    doc = ET.fromstring(data)
                    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                    cell = doc.find('.//s:c[@r="A3"]', ns); cell.set('t', kind)
                    cell.find('s:v', ns).text = value
                    data = ET.tostring(doc)
                target.writestr(item, data)
        return output

    def test_docx_examples_preserve_adjacent_run_formats_and_original(self):
        from docx import Document
        self.run_examples('office-docx')
        original, edited = Document(self.root/'report.docx'), Document(self.root/'edited.docx')
        self.assertEqual(original.paragraphs[1].text, 'Draft technical summary')
        self.assertEqual(edited.paragraphs[1].text, 'Final technical summary')
        self.assertTrue(edited.paragraphs[1].runs[0].bold)
        self.assertTrue(edited.paragraphs[1].runs[1].italic)
        self.assertEqual(len(edited.tables), 1)

    def test_xlsx_example_keeps_identifier_and_formula_without_inventing_cache(self):
        from openpyxl import load_workbook
        self.run_examples('office-xlsx')
        path = self.root/'report.xlsx'
        book = load_workbook(path, data_only=False)
        self.assertEqual(book['Revenue']['A2'].value, '001')
        self.assertEqual(book['Revenue']['B4'].data_type, 'f')
        self.assertEqual(book['Revenue']['B4'].value, '=SUM(B2:B3)')
        book.close()
        cached = load_workbook(path, data_only=True)
        self.assertIsNone(cached['Revenue']['B4'].value)
        cached.close()

    def test_pdf_input_needs_no_office_and_renders_only_selected_page(self):
        from pypdf import PdfWriter
        source = self.root/'input.pdf'; writer = PdfWriter()
        writer.add_blank_page(width=100, height=100); writer.add_blank_page(width=100, height=100)
        writer.write(source)
        original = source.read_bytes()
        def render(argv, directory, timeout):
            self.assertEqual(argv[1:5], ['-f', '2', '-l', '2'])
            Path(argv[-1]).with_suffix('.png').write_bytes(b'\x89PNG\r\n\x1a\nfixture')
        with patch.object(ops, 'executable', return_value='/fixture/pdftoppm') as lookup, patch.object(ops, 'run_command', side_effect=render):
            report = ops.perform('render', source, self.root/'rendered', pages='2')
        lookup.assert_called_once_with(None, ('pdftoppm',))
        self.assertEqual(report['rendered_pages'], [2])
        self.assertEqual(report['pdf_page_count'], 2)
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(report['visual_review'], 'not_run')

    def test_malformed_pdf_persists_failed_report(self):
        source = self.root/'bad.pdf'; source.write_bytes(b'%PDF-1.4\nnot parseable')
        with patch.object(ops, 'executable', return_value='/fixture/pdftoppm'):
            with self.assertRaises(Exception):
                ops.perform('render', source, self.root/'bad-render')
        self.assertEqual(json.loads((self.root/'bad-render/operation-report.json').read_text())['status'], 'failed')

    def test_recalculation_distinguishes_missing_error_and_empty_string_caches(self):
        source = self.workbook()
        self.assertEqual(ops.verify_recalculation(source, source)['missing_cache_cells'], ['Revenue!A3'])
        cached = self.cached_copy(source)
        self.assertEqual(ops.verify_recalculation(source, cached)['missing_cache_cells'], [])
        errored = self.cached_copy(source, '#DIV/0!', 'e')
        self.assertEqual(ops.verify_recalculation(source, errored)['formula_errors'], ['Revenue!A3: #DIV/0!'])
        blank = self.cached_copy(source, '', 'str')
        self.assertEqual(ops.verify_recalculation(source, blank)['missing_cache_cells'], [])

    def test_recalculation_rejects_formulas_replaced_by_constants(self):
        from openpyxl import load_workbook
        source = self.workbook(); result = self.root/'constant.xlsx'
        book = load_workbook(source); book['Revenue']['A3'] = 30; book.save(result); book.close()
        with self.assertRaisesRegex(ValueError, 'Formula locations or expressions changed'):
            ops.verify_recalculation(source, result)

    def test_recalculation_pipeline_preserves_source_and_output_formula_cache(self):
        from openpyxl import load_workbook
        source = self.workbook(); original = source.read_bytes(); cached = self.cached_copy(source)
        def convert(argv, directory, timeout):
            self.assertEqual(argv[argv.index('--convert-to')+1], 'xlsx:Calc MS Excel 2007 XML')
            (Path(argv[argv.index('--outdir')+1])/'source.xlsx').write_bytes(cached.read_bytes())
        with patch.object(ops, 'executable', return_value='/fixture/soffice'), patch.object(ops, 'run_command', side_effect=convert):
            report = ops.perform('recalculate', source, self.root/'recalc')
        self.assertEqual(report['status'], 'completed')
        result = load_workbook(self.root/'recalc/recalculated.xlsx', data_only=True)
        self.assertEqual(result['Revenue']['A3'].value, 30); result.close()
        self.assertEqual(source.read_bytes(), original)

    def test_resaving_stale_caches_without_calculating_cannot_pass(self):
        from openpyxl import load_workbook
        source = self.cached_copy(self.workbook(), '999')
        original = source.read_bytes()
        def no_calculation(argv, directory, timeout):
            staged = Path(argv[-1])
            book = load_workbook(staged, data_only=True)
            self.assertIsNone(book['Revenue']['A3'].value)
            book.close()
            (Path(argv[argv.index('--outdir')+1])/'source.xlsx').write_bytes(staged.read_bytes())
        with patch.object(ops, 'executable', return_value='/fixture/soffice'), patch.object(ops, 'run_command', side_effect=no_calculation):
            with self.assertRaisesRegex(ValueError, 'Formula caches are missing'):
                ops.perform('recalculate', source, self.root/'stale')
        report = json.loads((self.root/'stale/operation-report.json').read_text())
        self.assertEqual(report['status'], 'failed')
        self.assertEqual(report['recalculation']['missing_cache_cells'], ['Revenue!A3'])
        self.assertTrue((self.root/'stale/recalculated.xlsx').is_file())
        self.assertEqual(source.read_bytes(), original)

    def test_external_data_formula_fails_before_engine_or_output_creation(self):
        from openpyxl import load_workbook
        source = self.workbook(); book = load_workbook(source)
        book['Revenue']['A3'] = '=WEBSERVICE("https://example.invalid/")'; book.save(source); book.close()
        with patch.object(ops, 'executable', return_value='/fixture/soffice'), patch.object(ops, 'run_command') as command:
            with self.assertRaisesRegex(ValueError, 'External-data formulas'):
                ops.perform('recalculate', source, self.root/'blocked')
        command.assert_not_called()
        self.assertFalse((self.root/'blocked').exists())


if __name__ == '__main__':
    unittest.main()
