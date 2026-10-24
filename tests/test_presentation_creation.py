"""Exercise actual editable PPTX objects; visual rendering remains a separate integration check."""
import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT/'plugins/openas/skills/presentation-creation'
SCRIPTS = SKILL/'scripts'
sys.path.insert(0, str(SCRIPTS))
import deck_model as model
import build_deck as builder
import render_deck as renderer
sys.path.pop(0)
EXAMPLE = json.loads((SKILL/'assets/example-deck.json').read_text())
HAS_PPTX = importlib.util.find_spec('pptx') is not None and importlib.util.find_spec('PIL') is not None


class DeckModelTests(unittest.TestCase):
    def test_example_has_honest_time_budget_and_data_sources(self):
        model.validate(copy.deepcopy(EXAMPLE), SKILL/'assets')
        self.assertLessEqual(sum(x['seconds'] for x in EXAMPLE['slides']), EXAMPLE['duration_minutes']*60)
        chart = next(x for x in EXAMPLE['slides'] if x['type'] == 'chart')
        self.assertTrue(chart['sources'])
        self.assertIn('不代表真实', chart['caption'])

    def test_malformed_and_unknown_fields_are_rejected(self):
        invalid = [None, {}, {**EXAMPLE, 'version': True}, {**EXAMPLE, 'slides': ['invalid']},
                   {**EXAMPLE, 'theme': []}, {**EXAMPLE, 'template': 'unsupported.pptx'},
                   {**EXAMPLE, 'theme': {'accent': '#FF0000'}}, {**EXAMPLE, 'theme': {'unknown': 'x'}}]
        for value in invalid:
            with self.subTest(value=value), self.assertRaises(ValueError):
                model.validate(value, SKILL/'assets')
        for changes in ({'sources': 'url'}, {'notes': ''}, {'seconds': 0}, {'type': 'unknown'}, {'subtitle': 'Not supported here'}):
            spec = copy.deepcopy(EXAMPLE)
            spec['slides'] = [{**spec['slides'][1], **changes}]
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                model.validate(spec, SKILL/'assets')

    def test_chart_requires_aligned_finite_values_units_and_sources(self):
        chart = copy.deepcopy(EXAMPLE['slides'][4])
        invalid = [dict(chart, series=[{'name': 'bad', 'values': [1]}]),
                   dict(chart, series=[{'name': 'bad', 'values': [1, 2, float('nan')]}]),
                   dict(chart, sources=[]), dict(chart, unit=''), dict(chart, chart_type='3d')]
        for value in invalid:
            with self.subTest(value=value), self.assertRaises(ValueError):
                model.validate({**EXAMPLE, 'slides': [value]}, SKILL/'assets')

    def test_timing_limit_rejects_overbooked_deck(self):
        with self.assertRaises(ValueError):
            model.validate({**EXAMPLE, 'duration_minutes': .1}, SKILL/'assets')

    def test_image_assets_cannot_escape_project_root(self):
        with tempfile.TemporaryDirectory() as temporary:
            base = Path(temporary); assets = base/'assets'; assets.mkdir()
            outside = base/'outside.png'; outside.write_bytes(b'fixture')
            (assets/'link.png').symlink_to(outside)
            for name in ('../outside.png', str(outside), 'link.png'):
                with self.subTest(name=name), self.assertRaises(ValueError):
                    model.image_path(assets, name)

    def test_missing_renderer_does_not_install_tools_or_create_output(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)/'preview'
            with patch.object(renderer.shutil, 'which', return_value=None):
                with self.assertRaises(ValueError):
                    renderer.render(Path(temporary)/'deck.pptx', output)
            self.assertFalse(output.exists())


@unittest.skipUnless(HAS_PPTX, 'Install presentation-creation/requirements.txt to test native PPTX generation')
class EditableDeckTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(); self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def write(self, spec=None, name='deck.pptx'):
        path = self.root/name
        path.write_bytes(builder.build(copy.deepcopy(spec or EXAMPLE), self.root))
        return path

    def test_generated_deck_preserves_text_notes_native_chart_and_table(self):
        from pptx import Presentation
        deck = self.write()
        report = model.inspect_pptx(deck)
        self.assertTrue(report['passed'], report['errors'])
        self.assertEqual(report['slide_count'], 8)
        self.assertTrue(any(x['native_chart'] for s in report['slides'] for x in s['shapes']))
        self.assertTrue(any(x['native_table'] for s in report['slides'] for x in s['shapes']))
        prs = Presentation(deck)
        self.assertEqual(prs.core_properties.author, '')
        self.assertEqual(prs.core_properties.last_modified_by, 'OpenAS')
        self.assertIn(EXAMPLE['slides'][0]['notes'], prs.slides[0].notes_slide.notes_text_frame.text)
        chart = next(s.chart for s in prs.slides[4].shapes if s.has_chart)
        self.assertEqual(list(chart.series[0].values), [3.0, 5.0, 4.0])
        self.assertFalse(chart.has_title)
        table = next(s.table for s in prs.slides[5].shapes if s.has_table)
        self.assertEqual(table.cell(1, 0).text, 'PPTX')
        with zipfile.ZipFile(deck) as archive:
            self.assertTrue(any(x.startswith('ppt/embeddings/') and x.endswith('.xlsx') for x in archive.namelist()))
            xml = archive.read('ppt/slides/slide1.xml').decode()
            self.assertIn('Microsoft YaHei', xml)
            self.assertIn('Arial', xml)

    def test_standard_aspect_ratio_and_line_chart_are_native(self):
        from pptx import Presentation
        spec = copy.deepcopy(EXAMPLE)
        spec['size'] = 'standard'
        chart = spec['slides'][4]; chart['chart_type'] = 'line'
        spec['slides'] = [chart]
        path = self.write(spec)
        report = model.inspect_pptx(path)
        self.assertTrue(report['passed'])
        self.assertAlmostEqual(report['width']/report['height'], 4/3, places=5)
        prs = Presentation(path)
        chart = next(x.chart for x in prs.slides[0].shapes if x.has_chart)
        self.assertEqual(list(chart.series[0].values), [3.0, 5.0, 4.0])

    def test_image_is_original_aspect_preserved_and_alt_text_recorded(self):
        from PIL import Image
        from pptx import Presentation
        Image.new('RGB', (200, 100), 'teal').save(self.root/'figure.png')
        slide = {'type': 'image', 'title': 'Original figure', 'image': 'figure.png', 'alt': 'A two-to-one fixture image',
                 'notes': 'Explain original figure', 'seconds': 30, 'caption': 'Original source figure'}
        path = self.write({**EXAMPLE, 'slides': [slide]})
        prs = Presentation(path)
        image = next(x for x in prs.slides[0].shapes if x.name == 'original-image')
        self.assertAlmostEqual(image.width/image.height, 2)
        self.assertEqual(image._element.nvPicPr.cNvPr.get('descr'), slide['alt'])
        self.assertTrue(model.inspect_pptx(path)['passed'])

    def test_geometry_checker_detects_out_of_bounds_and_overlap_in_actual_xml(self):
        from pptx import Presentation
        from pptx.util import Inches
        path = self.write()
        prs = Presentation(path)
        shape = prs.slides[1].shapes.add_textbox(Inches(20), Inches(2), Inches(2), Inches(1))
        shape.text = 'Outside canvas'
        first = prs.slides[1].shapes[2]
        # Use two content boxes rather than a decorative marker.
        extra = prs.slides[1].shapes.add_textbox(first.left, first.top, Inches(2), Inches(1))
        extra.text = 'Overlapping content'
        prs.save(path)
        report = model.inspect_pptx(path)
        self.assertFalse(report['passed'])
        self.assertTrue(any('out-of-bounds' in x for x in report['errors']))
        self.assertTrue(any('overlapping content' in x for x in report['errors']))

    def test_excessive_text_fails_instead_of_shrinking_to_unreadable_size(self):
        spec = copy.deepcopy(EXAMPLE)
        spec['slides'] = [{**spec['slides'][1], 'bullets': ['非常长的正文'*40]*6}]
        with self.assertRaises(ValueError):
            builder.build(spec, self.root)

    def test_cli_refuses_existing_output_and_preserves_original_bytes(self):
        path = self.root/'existing.pptx'; path.write_bytes(b'original-user-file')
        process = subprocess.run([sys.executable, str(SCRIPTS/'build_deck.py'), '--spec', str(SKILL/'assets/example-deck.json'),
                                  '--output', str(path)], capture_output=True, text=True)
        self.assertEqual(process.returncode, 2)
        self.assertEqual(path.read_bytes(), b'original-user-file')

    def test_cli_reports_real_generation_and_visual_review_not_run(self):
        path = self.root/'new.pptx'
        process = subprocess.run([sys.executable, str(SCRIPTS/'build_deck.py'), '--spec', str(SKILL/'assets/example-deck.json'),
                                  '--output', str(path)], capture_output=True, text=True)
        self.assertEqual(process.returncode, 0, process.stderr)
        report = json.loads(process.stdout)
        self.assertTrue(report['structure_passed'])
        self.assertEqual(report['rendered_review'], 'not_run')
        self.assertTrue(zipfile.is_zipfile(path))

    def test_mock_renderer_keeps_visual_review_unverified(self):
        from PIL import Image
        deck = self.write({**EXAMPLE, 'slides': EXAMPLE['slides'][:2]})
        output = self.root/'preview'

        def fake_execute(command, directory, timeout):
            if '--convert-to' in command:
                (directory/(deck.stem+'.pdf')).write_bytes(b'%PDF-fixture')
            else:
                for index in (1, 2):
                    Image.new('RGB', (1600, 900), '#FFFFFF').save(directory/f'pages/slide-{index}.png')

        with patch.object(renderer.shutil, 'which', return_value='/fixture/renderer'), \
                patch.object(renderer, 'execute', side_effect=fake_execute):
            result = renderer.render(deck, output)
        self.assertEqual(len(result['pages']), 2)
        self.assertEqual(result['human_visual_review'], 'not_run')
        self.assertTrue((output/'montage.png').is_file())

    def test_partial_pdf_export_does_not_claim_rendering_completed(self):
        from PIL import Image
        deck = self.write()
        output = self.root/'partial-preview'

        def fake_execute(command, directory, timeout):
            if '--convert-to' in command:
                (directory/(deck.stem+'.pdf')).write_bytes(b'%PDF-fixture')
            else:
                Image.new('RGB', (1600, 900), '#FFFFFF').save(directory/'pages/slide-1.png')

        with patch.object(renderer.shutil, 'which', return_value='/fixture/renderer'), \
                patch.object(renderer, 'execute', side_effect=fake_execute):
            with self.assertRaisesRegex(ValueError, 'page count differs'):
                renderer.render(deck, output)
        self.assertFalse((output/'render-report.json').exists())


if __name__ == '__main__':
    unittest.main()
